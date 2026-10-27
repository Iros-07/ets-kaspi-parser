import hashlib
import io
import re
from urllib.parse import urljoin

import psycopg2
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from db_config import DB

HEADERS = {"User-Agent": "Mozilla/5.0"}
ETS_URL = "https://ets.kz/documents/trading-results/"
KASPI_URL = "https://ccx.kz/trading-and-analytics/trading-results?lang=ru"


def connect_server():
    cfg = {**DB, "dbname": "postgres"}
    return psycopg2.connect(**cfg)


def ensure_database():
    conn = connect_server()
    conn.autocommit = True
    cur = conn.cursor()
    cur.execute("SELECT 1 FROM pg_database WHERE datname = %s", (DB["dbname"],))
    if not cur.fetchone():
        cur.execute(f'CREATE DATABASE "{DB["dbname"]}"')
        print("создана БД:", DB["dbname"])
    cur.close()
    conn.close()


def connect():
    ensure_database()
    conn = psycopg2.connect(**DB)
    cur = conn.cursor()
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS trades (
            id SERIAL PRIMARY KEY,
            exchange TEXT NOT NULL,
            trade_date DATE NOT NULL,
            product_name TEXT,
            tnved TEXT,
            product_code TEXT,
            buyer_name TEXT,
            buyer_bin TEXT,
            buyer_broker TEXT,
            seller_name TEXT,
            seller_bin TEXT,
            seller_broker TEXT,
            deals_count INTEGER,
            price_avg NUMERIC,
            amount NUMERIC,
            delivery TEXT,
            uniq_key TEXT NOT NULL UNIQUE,
            created_at TIMESTAMP DEFAULT NOW()
        )
        """
    )
    cur.execute(
        "CREATE INDEX IF NOT EXISTS idx_trades_date ON trades(trade_date)"
    )
    cur.execute(
        "CREATE INDEX IF NOT EXISTS idx_trades_exchange ON trades(exchange)"
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS sync_log (
            id SERIAL PRIMARY KEY,
            started_at TIMESTAMP NOT NULL,
            finished_at TIMESTAMP,
            exchange TEXT NOT NULL,
            fetched INTEGER DEFAULT 0,
            added INTEGER DEFAULT 0,
            status TEXT NOT NULL,
            error_text TEXT
        )
        """
    )
    conn.commit()
    cur.close()
    return conn


def to_iso(date_str):
    # дд.мм.гггг -> гггг-мм-дд
    m = re.match(r"(\d{2})\.(\d{2})\.(\d{4})", str(date_str).strip())
    if not m:
        return str(date_str)
    d, mo, y = m.groups()
    return f"{y}-{mo}-{d}"


def clean_text(v):
    if v is None:
        return None
    s = " ".join(str(v).split()).strip()
    return s or None


def clean_num(v):
    if v is None or v == "" or v == "-":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(" ", "").replace("\xa0", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def split_name_bin(text):
    # имя и бин из одной строки
    text = clean_text(text)
    if not text:
        return None, None
    m = re.search(r"^(.*?)(?:\s+)(\d{12})$", text)
    if m:
        return clean_text(m.group(1)), m.group(2)
    return text, None


def make_key(row):
    parts = [
        row["exchange"],
        str(row["trade_date"]),
        row.get("product_name") or "",
        row.get("product_code") or "",
        row.get("buyer_name") or "",
        row.get("buyer_bin") or "",
        row.get("seller_name") or "",
        row.get("seller_bin") or "",
        str(row.get("price_avg") or ""),
        str(row.get("amount") or ""),
        row.get("delivery") or "",
    ]
    raw = "|".join(parts).lower()
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()


def save_rows(conn, rows):
    sql = """
        INSERT INTO trades (
            exchange, trade_date, product_name, tnved, product_code,
            buyer_name, buyer_bin, buyer_broker,
            seller_name, seller_bin, seller_broker,
            deals_count, price_avg, amount, delivery, uniq_key
        ) VALUES (
            %(exchange)s, %(trade_date)s, %(product_name)s, %(tnved)s, %(product_code)s,
            %(buyer_name)s, %(buyer_bin)s, %(buyer_broker)s,
            %(seller_name)s, %(seller_bin)s, %(seller_broker)s,
            %(deals_count)s, %(price_avg)s, %(amount)s, %(delivery)s, %(uniq_key)s
        )
        ON CONFLICT (uniq_key) DO NOTHING
    """
    added = 0
    cur = conn.cursor()
    for row in rows:
        row["uniq_key"] = make_key(row)
        cur.execute(sql, row)
        added += cur.rowcount
    conn.commit()
    cur.close()
    return added


def get_html(url):
    r = requests.get(url, headers=HEADERS, timeout=30)
    r.raise_for_status()
    return BeautifulSoup(r.text, "lxml")


def fetch_ets():
    print("ETS: ищу последний файл...")
    soup = get_html(ETS_URL)
    year_url = None
    for a in soup.find_all("a", href=True):
        t = a.get_text(" ", strip=True)
        if re.search(r"Итоги торгов\s+\d{4}\s+год", t, re.I):
            year_url = urljoin(ETS_URL, a["href"])

    soup = get_html(year_url)
    months = []
    for a in soup.find_all("a", href=True):
        t = a.get_text(" ", strip=True)
        m = re.search(r"Итоги торгов\s+\d{4}\s+год\s+(\S+)", t, re.I)
        if m:
            months.append((m.group(1), urljoin(year_url, a["href"])))
    month_name, month_url = months[-1]

    soup = get_html(month_url)
    files = []
    seen = set()
    for a in soup.find_all("a", href=True):
        href = a["href"]
        if ".xls" not in href.lower():
            continue
        link = urljoin(month_url, href)
        if link in seen:
            continue
        seen.add(link)
        parent = a.find_parent(["li", "div", "p", "td"])
        text = parent.get_text(" ", strip=True) if parent else ""
        m = re.search(r"(\d{2}\.\d{2}\.\d{4})", text)
        date = m.group(1) if m else "?"
        files.append((date, link))

    date, file_url = files[-1]
    print("ETS:", month_name, date, file_url)

    r = requests.get(file_url, headers=HEADERS, timeout=60)
    wb = load_workbook(io.BytesIO(r.content), data_only=True)
    ws = wb.active

    out = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        product = row[7]
        if product is None:
            continue
        if str(product).strip().lower() == "итого":
            continue

        out.append({
            "exchange": "ETS",
            "trade_date": to_iso(date),
            "product_name": clean_text(product),
            "tnved": clean_text(row[8]),
            "product_code": clean_text(row[9]),
            "buyer_name": clean_text(row[1]),
            "buyer_bin": clean_text(row[2]),
            "buyer_broker": clean_text(row[3]),
            "seller_name": clean_text(row[4]),
            "seller_bin": clean_text(row[5]),
            "seller_broker": clean_text(row[6]),
            "deals_count": int(row[10]) if row[10] is not None else None,
            "price_avg": clean_num(row[15]),
            "amount": clean_num(row[16]),
            "delivery": None,
        })
    return out


def fetch_kaspi():
    print("Каспий: читаю таблицу...")
    r = requests.get(KASPI_URL, headers=HEADERS, timeout=30)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "lxml")
    table = soup.find("table")
    if not table:
        print("Каспий: таблица не найдена")
        return []

    out = []
    for tr in table.find_all("tr")[1:]:
        cols = [td.get_text(" ", strip=True) for td in tr.find_all("td")]
        if len(cols) < 16:
            continue
        product = cols[3]
        if "итого" in product.lower():
            continue

        buyer_name, buyer_bin = split_name_bin(cols[11])
        seller_name, seller_bin = split_name_bin(cols[13])

        out.append({
            "exchange": "Каспий",
            "trade_date": to_iso(cols[0]),
            "product_name": clean_text(product),
            "tnved": clean_text(cols[2] if cols[2] != "-" else None),
            "product_code": None,
            "buyer_name": buyer_name,
            "buyer_bin": buyer_bin,
            "buyer_broker": clean_text(cols[10]),
            "seller_name": seller_name,
            "seller_bin": seller_bin,
            "seller_broker": clean_text(cols[12]),
            "deals_count": int(clean_num(cols[14]) or 0) or None,
            "price_avg": clean_num(cols[8]),
            "amount": clean_num(cols[9]),
            "delivery": clean_text(cols[15] if cols[15] != "-" else None),
        })
    return out


def log_run_ok(conn, exchange, started_at, fetched, added):
    cur = conn.cursor()
    cur.execute(
        """
        INSERT INTO sync_log (
            started_at, finished_at, exchange, fetched, added, status, error_text
        ) VALUES (%s, NOW(), %s, %s, %s, 'ok', NULL)
        """,
        (started_at, exchange, fetched, added),
    )
    conn.commit()
    cur.close()


def log_run_err(conn, exchange, started_at, error_text):
    cur = conn.cursor()
    cur.execute(
        """
        INSERT INTO sync_log (
            started_at, finished_at, exchange, fetched, added, status, error_text
        ) VALUES (%s, NOW(), %s, 0, 0, 'error', %s)
        """,
        (started_at, exchange, error_text[:2000] if error_text else None),
    )
    conn.commit()
    cur.close()


def main():
    from datetime import datetime

    conn = connect()
    print("PostgreSQL:", DB["host"], DB["dbname"])

    started = datetime.now()
    try:
        ets_rows = fetch_ets()
        n1 = save_rows(conn, ets_rows)
        log_run_ok(conn, "ETS", started, len(ets_rows), n1)
    except Exception as e:
        print("ETS ошибка:", e)
        log_run_err(conn, "ETS", started, str(e))
        ets_rows, n1 = [], 0

    started = datetime.now()
    try:
        kaspi_rows = fetch_kaspi()
        n2 = save_rows(conn, kaspi_rows)
        log_run_ok(conn, "Каспий", started, len(kaspi_rows), n2)
    except Exception as e:
        print("Каспий ошибка:", e)
        log_run_err(conn, "Каспий", started, str(e))
        kaspi_rows, n2 = [], 0

    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM trades")
    total = cur.fetchone()[0]
    cur.execute("SELECT exchange, COUNT(*) FROM trades GROUP BY exchange")
    by_ex = dict(cur.fetchall())
    cur.close()
    conn.close()

    print()
    print(f"ETS: получено {len(ets_rows)}, новых {n1}")
    print(f"Каспий: получено {len(kaspi_rows)}, новых {n2}")
    print(f"всего в БД: {total}", by_ex)


if __name__ == "__main__":
    main()
