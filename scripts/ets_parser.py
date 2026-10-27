# парсер ETS - качает последний excel и печатает в терминал
import io
import re
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from tabulate import tabulate

URL = "https://ets.kz/documents/trading-results/"
HEADERS = {"User-Agent": "Mozilla/5.0"}


def poluchit_html(url):
    otvet = requests.get(url, headers=HEADERS, timeout=30)
    otvet.raise_for_status()
    return BeautifulSoup(otvet.text, "lxml")


# 1) находим ссылку на год
soup = poluchit_html(URL)
year_url = None

for a in soup.find_all("a", href=True):
    tekst = a.get_text(" ", strip=True)
    if "Итоги торгов" in tekst and "год" in tekst:
        # ссылки с месяцем пропускаем тут грубо
        if "Октябрь" in tekst or "октябрь" in tekst:
            continue
        if "Январь" in tekst or "Февраль" in tekst:
            continue
        year_url = urljoin(URL, a["href"])

print("год:", year_url)

# 2) месяцы
soup = poluchit_html(year_url)
mesyacy = []

for a in soup.find_all("a", href=True):
    tekst = a.get_text(" ", strip=True)
    if "Итоги торгов" not in tekst:
        continue
    if tekst.strip().endswith("год"):
        continue

    chasti = tekst.split(" ")
    if len(chasti) < 2:
        continue

    imya_mesyaca = chasti[-1]
    ssylka = urljoin(year_url, a["href"])
    mesyacy.append([imya_mesyaca, ssylka])

posledniy_mesyac = mesyacy[len(mesyacy) - 1]
month_name = posledniy_mesyac[0]
month_url = posledniy_mesyac[1]
print("месяц:", month_name, month_url)

# 3) excel файлы
soup = poluchit_html(month_url)
fayly = []
uzhe = []

for a in soup.find_all("a", href=True):
    href = a["href"]
    if ".xls" not in href.lower():
        continue

    polnaya = urljoin(month_url, href)
    if polnaya in uzhe:
        continue
    uzhe.append(polnaya)

    roditel = a.find_parent("li")
    if roditel is None:
        roditel = a.find_parent("div")

    if roditel is not None:
        podpis = roditel.get_text(" ", strip=True)
    else:
        podpis = ""

    data = "?"
    naydeno = re.search(r"(\d{2}\.\d{2}\.\d{4})", podpis)
    if naydeno:
        data = naydeno.group(1)

    fayly.append([data, polnaya])

posledniy = fayly[len(fayly) - 1]
data = posledniy[0]
file_url = posledniy[1]
print("файл:", data, file_url)

# 4) читаем excel
otvet = requests.get(file_url, headers=HEADERS, timeout=60)
kniga = load_workbook(io.BytesIO(otvet.content), data_only=True)
list_excel = kniga.active

rows = []
for row in list_excel.iter_rows(min_row=5, values_only=True):
    pokupatel = row[1]
    prodavec = row[4]
    tovar = row[7]
    cena = row[15]
    summa = row[16]

    if tovar is None:
        continue

    if str(tovar).strip().lower() == "итого":
        continue

    p = str(tovar)
    if len(p) > 48:
        p = p[0:40] + "..."

    b = str(pokupatel or "")
    if len(b) > 28:
        b = b[0:20] + "..."

    s = str(prodavec or "")
    if len(s) > 48:
        s = s[0:40] + "..."

    rows.append([data, p, b, s, cena, summa])

print()
print("=== ETS", data, "===")
print()
print(
    tabulate(
        rows,
        headers=["Дата", "Товар", "Покупатель", "Продавец", "Цена ср.", "Сумма"],
        tablefmt="github",
    )
)
print()
print("Строк:", len(rows))
